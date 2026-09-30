#!/usr/bin/env python3
"""Extract per-player projections from the projection workbook into a compact JSON file.

The workbook is the kitchen's private input. Its numbers are used to ORDER players and to place
tier breaks; they are never published. This script only reads cached cell values (data_only=True),
so the workbook must have been saved by Excel or recalculated (LibreOffice) before upload.

Usage
  python3 tools/extract_projections.py WORKBOOK.xlsx --out projections.json
  python3 tools/extract_projections.py WORKBOOK.xlsx --inspect            # list sheets and candidate header rows
  python3 tools/extract_projections.py WORKBOOK.xlsx --sheet "Summary" --map player=Player,pos=Pos,team=Tm,ppg=PPR/G

Discovery: every sheet is scanned for a header row that contains a player-name column and a
position column. Rows below it are read until a run of blank rows. All numeric columns are kept
under their header text, and a few normalized fields are filled when a header matches:
  player, team, pos, games, points (season total), ppg (points per game), week_pts (this week's projection)
The mapping in tools/workbook_map.json (if present) is applied first and overrides discovery; use it
once the workbook's layout is known so every run reads the same cells.
"""
import argparse, json, os, re, sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl --break-system-packages")

HERE = os.path.dirname(os.path.abspath(__file__))
MAP_FILE = os.path.join(HERE, "workbook_map.json")

PLAYER_HEADERS = ["player", "name", "player name"]
POS_HEADERS = ["pos", "position"]
TEAM_HEADERS = ["team", "tm", "nfl team"]
GAMES_HEADERS = ["g", "gp", "games", "gms"]
POINTS_HEADERS = ["fpts", "fantasy points", "points", "pts", "ppr", "ppr pts", "total", "proj", "projection", "fp"]
PPG_HEADERS = ["ppg", "pts/g", "fpts/g", "ppr/g", "points/g", "pts per game", "fp/g", "fppg"]
WEEK_HEADERS = ["week", "wk", "this week", "week pts", "wk pts"]
POS_VALUES = {"QB", "RB", "WR", "TE", "K", "DST", "D/ST", "DEF", "FB"}

def norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip().lower())

def header_index(headers, candidates):
    for i, h in enumerate(headers):
        if h in candidates:
            return i
    for i, h in enumerate(headers):
        for c in candidates:
            if h.startswith(c + " ") or h.endswith(" " + c):
                return i
    return None

def find_tables(ws, max_scan_rows=60):
    """Yield (header_row_index, headers) for rows that look like a player table header."""
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_scan_rows), values_only=True))
    for ri, row in enumerate(rows):
        headers = [norm(c) for c in row]
        if header_index(headers, PLAYER_HEADERS) is not None and header_index(headers, POS_HEADERS) is not None:
            yield ri + 1, headers
        elif header_index(headers, PLAYER_HEADERS) is not None:
            # a table whose position is implied by the sheet (e.g. a sheet named "RB")
            if norm(ws.title).upper() in POS_VALUES or re.search(r"\b(qb|rb|wr|te)s?\b", norm(ws.title)):
                yield ri + 1, headers


POSRANK_RE = re.compile(r"^(QB|RB|WR|TE|K|DST)\s*(\d+)$", re.I)
PLACEHOLDER_RE = re.compile(r"^[A-Z]{2,3}\s+(QB|RB|WR|TE|K|DST)\d+$")

def team_lookup(wb, sheet_names, aliases):
    """Map (player, pos) and player -> team code from per-team sheets whose tables carry a Player
    column and a Positional Rank column (values like RB8). The team is the sheet name."""
    by_key, by_name = {}, {}
    for name in sheet_names:
        if name not in wb.sheetnames:
            print(f"warning: team sheet '{name}' not found", file=sys.stderr); continue
        team = aliases.get(name.strip().upper(), name.strip().upper())
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        for ri, row in enumerate(rows):
            headers = [norm(c) for c in row]
            ip = header_index(headers, PLAYER_HEADERS)
            ir = next((i for i, h in enumerate(headers) if h.startswith("positional rank")), None)
            if ip is None or ir is None:
                continue
            blanks = 0
            for r in rows[ri + 1:]:
                nm = r[ip] if ip < len(r) else None
                if nm is None or str(nm).strip() == "":
                    blanks += 1
                    if blanks >= 2:
                        break
                    continue
                blanks = 0
                if norm(nm) == "player":
                    break
                nm = str(nm).strip()
                if PLACEHOLDER_RE.match(nm):
                    continue
                pr = str(r[ir] or "").strip() if ir < len(r) else ""
                m = POSRANK_RE.match(pr)
                if not m:
                    continue
                pos = m.group(1).upper()
                by_key[(nm.lower(), pos)] = team
                by_name.setdefault(nm.lower(), set()).add(team)
    return by_key, {k: next(iter(v)) for k, v in by_name.items() if len(v) == 1}

def read_table(ws, header_row, headers, sheet_pos=None, mapping=None):
    mapping = mapping or {}
    def col(key, cands):
        if key in mapping:
            m = norm(mapping[key])
            return headers.index(m) if m in headers else None
        return header_index(headers, cands)
    ip = col("player", PLAYER_HEADERS); ipos = col("pos", POS_HEADERS); it = col("team", TEAM_HEADERS)
    ig = col("games", GAMES_HEADERS); ipt = col("points", POINTS_HEADERS); ippg = col("ppg", PPG_HEADERS); iw = col("week_pts", WEEK_HEADERS)
    out, blanks = [], 0
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        name = row[ip] if ip is not None and ip < len(row) else None
        if name is None or str(name).strip() == "":
            blanks += 1
            if blanks >= 3:
                break
            continue
        blanks = 0
        if PLACEHOLDER_RE.match(str(name).strip()):
            continue
        rec = {"player": str(name).strip(), "sheet": ws.title}
        if ipos is not None and ipos < len(row) and row[ipos]:
            rec["pos"] = str(row[ipos]).strip().upper()
        elif sheet_pos:
            rec["pos"] = sheet_pos
        if it is not None and it < len(row) and row[it]:
            rec["team"] = str(row[it]).strip().upper()
        nums = {}
        for i, h in enumerate(headers):
            if i < len(row) and isinstance(row[i], (int, float)) and h and i not in (ip, ipos, it):
                nums[h] = round(float(row[i]), 3)
        rec["values"] = nums
        for key, idx in (("games", ig), ("points", ipt), ("ppg", ippg), ("week_pts", iw)):
            if idx is not None and idx < len(row) and isinstance(row[idx], (int, float)):
                rec[key] = round(float(row[idx]), 3)
        if "ppg" not in rec and "points" in rec and rec.get("games"):
            rec["ppg"] = round(rec["points"] / rec["games"], 3)
        out.append(rec)
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--out", default="projections.json")
    ap.add_argument("--inspect", action="store_true")
    ap.add_argument("--sheet", help="only read this sheet (name)")
    ap.add_argument("--map", help="comma list key=Header for player,pos,team,games,points,ppg,week_pts")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.workbook, data_only=True, read_only=True)
    mapping = {}
    if os.path.exists(MAP_FILE):
        mapping = json.load(open(MAP_FILE))
    if a.map:
        for kv in a.map.split(","):
            k, v = kv.split("=", 1); mapping.setdefault("columns", {})[k.strip()] = v.strip()
    sheets = [a.sheet] if a.sheet else (mapping.get("sheets") or wb.sheetnames)
    cols = mapping.get("columns", {})

    if a.inspect:
        for name in wb.sheetnames:
            ws = wb[name]
            print(f"== {name}  ({ws.max_row} rows x {ws.max_column} cols)")
            for hr, headers in find_tables(ws):
                print(f"   header row {hr}: {[h for h in headers if h][:20]}")
        return

    players = []
    for name in sheets:
        if name not in wb.sheetnames:
            print(f"warning: sheet '{name}' not found", file=sys.stderr); continue
        ws = wb[name]
        sheet_pos = name.strip().upper() if name.strip().upper() in POS_VALUES else None
        if sheet_pos is None:
            m = re.search(r"\b(qb|rb|wr|te|k|dst)\b", norm(name))
            sheet_pos = m.group(1).upper() if m else None
        for hr, headers in find_tables(ws):
            players.extend(read_table(ws, hr, headers, sheet_pos, cols))
    # de-duplicate: keep the record with the most numeric values per (player, team, pos)
    best = {}
    for p in players:
        key = (p["player"].lower(), p.get("team", ""), p.get("pos", ""))
        if key not in best or len(p["values"]) > len(best[key]["values"]):
            best[key] = p
    players = sorted(best.values(), key=lambda p: (p.get("pos", ""), -(p.get("ppg") or p.get("points") or 0)))
    if not players:
        sys.exit("could not find a player table; run with --inspect and pass --sheet/--map or fill tools/workbook_map.json")
    # team from per-team sheets (mapping "team_sheets"), when the ranking tables carry none
    if mapping.get("team_sheets"):
        by_key, by_name = team_lookup(wb, mapping["team_sheets"], {k.upper(): v for k, v in (mapping.get("team_aliases") or {}).items()})
        for p in players:
            if not p.get("team"):
                t = by_key.get((p["player"].lower(), p.get("pos", ""))) or by_name.get(p["player"].lower())
                if t:
                    p["team"] = t
    # per-game number when the workbook has no games column (season projection / season_games)
    season_games = mapping.get("season_games")
    ppg_basis = None
    if season_games:
        for p in players:
            if "ppg" not in p and "points" in p:
                p["ppg"] = round(p["points"] / season_games, 3)
        ppg_basis = f"season points / {season_games} (no games column in the workbook)"
    meta = {"source": os.path.basename(a.workbook), "count": len(players), "players": players}
    if ppg_basis:
        meta["ppg_basis"] = ppg_basis
    with open(a.out, "w") as f:
        json.dump(meta, f, indent=1)
    no_team = [p["player"] for p in players if not p.get("team")]
    if no_team:
        print(f"warning: {len(no_team)} players have no team: " + ", ".join(no_team[:12]) + (" ..." if len(no_team) > 12 else ""), file=sys.stderr)
    by_pos = {}
    for p in players:
        by_pos[p.get("pos", "?")] = by_pos.get(p.get("pos", "?"), 0) + 1
    print(f"extracted {len(players)} players: " + ", ".join(f"{k} {v}" for k, v in sorted(by_pos.items())))
    missing = sum(1 for p in players if "ppg" not in p and "points" not in p)
    if missing:
        print(f"warning: {missing} players have no points column; check --map", file=sys.stderr)

if __name__ == "__main__":
    main()
